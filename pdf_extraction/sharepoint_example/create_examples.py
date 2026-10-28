#!/usr/bin/env python3
"""Create illustrative SharePoint list data; no tenant operations are made."""
import csv
import json
from pathlib import Path
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

ROOT=Path(__file__).resolve().parent


def col(name,label,kind='Text',note=''):
    return {'name':name,'display_name':label,'recommended_sharepoint_type':kind,'note':note}


def main():
    payload=json.loads((Path(__file__).resolve().parents[1] / 'order_mapping_example' / 'deliverables' / 'pa_runtime_payload.json').read_text(encoding='utf-8'))
    req=payload['requested_specification']
    lists=[{
        'list_name':'ExtractedItems','display_name':'문서별 추출값','grain':'발주서의 품목 한 줄당 한 행',
        'columns':[
            col('Title','문서 품목 키','Text','고유값 권장. 발주번호 + 품목순번으로 식별'),
            col('PurchaseOrder','발주번호','Text','숫자로 보이더라도 식별자는 텍스트로 저장'),
            col('ItemNo','품목순번','Number'),col('Department','부서','Choice','구매팀은 예시 부서명'),
            col('ProductFamily','제품군','Choice'),col('Standard','적용 규격'),col('SourceGrade','요구 강종'),
            col('SourceCoating','요구 도금'),col('ThicknessMm','두께(mm)','Number','단위는 mm. 최소 소수 셋째 자리까지 표시'),
            col('WidthMm','폭(mm)','Number'),col('ThicknessPlusMm','두께 상향 공차(mm)','Number'),
            col('ThicknessMinusMm','두께 하향 공차(mm)','Number'),col('CoilInnerDiameterMm','코일 내경(mm)','Number'),
            col('CoilWeightMinKg','코일 최소중량(kg)','Number'),col('CoilWeightMaxKg','코일 최대중량(kg)','Number'),
            col('QuantityT','주문수량(t)','Number'),col('SkinPassed','스킨패스 요구','Yes/No'),
            col('TensionLeveled','텐션레벨링 요구','Yes/No'),col('Chromate','크로메이트 요구'),
            col('HyundaiGrade','확정 현대제철 강종','Text','미확정이면 빈칸. 원래 요구 강종과 별도 관리'),
            col('HyundaiCoating','확정 현대제철 도금','Text','미확정이면 빈칸. 임의로 200이나 Z180을 넣지 않음'),
            col('Status','처리 상태','Choice','NeedsConfirmation / Ready / Completed / Failed'),
            col('TemplateKey','적용 템플릿 키','Text','PUR_SPEC는 예시 템플릿'),
            col('RuleSetVersion','적용 기준 버전','Text','처리에 사용한 설정 버전을 기록')],
        'rows':[{'Title':'PODEMO001-001','PurchaseOrder':'DEMO001','ItemNo':1,'Department':'구매팀',
                 'ProductFamily':'GI','Standard':req['standard'],'SourceGrade':req['grade'],'SourceCoating':req['coating_designation'],
                 'ThicknessMm':req['thickness_mm'],'WidthMm':req['width_mm'],
                 'ThicknessPlusMm':req['thickness_tolerance']['plus_mm'],'ThicknessMinusMm':0,
                 'CoilInnerDiameterMm':508,'CoilWeightMinKg':4000,'CoilWeightMaxKg':5000,'QuantityT':60,
                 'SkinPassed':False,'TensionLeveled':True,'Chromate':'CR6+',
                 'HyundaiGrade':None,'HyundaiCoating':None,'Status':'NeedsConfirmation',
                 'TemplateKey':'PUR_SPEC','RuleSetVersion':'EXAMPLE_V1'}]
    },{
        'list_name':'MappingRules','display_name':'수치·대응 기준','grain':'한 적용 조건과 대응값당 한 행',
        'columns':[
            col('Title','규칙 키'),col('RuleType','규칙 종류','Choice','Product / Grade / Coating'),
            col('ProductFamily','제품군'),col('SourceStandard','원본 적용 규격'),
            col('SourceValue','원본 값'),col('TargetValue','대응 값','Text','미확정이면 실제 빈칸'),
            col('Status','검토 상태','Choice','Confirmed / NeedsConfirmation'),
            col('IsActive','적용 여부','Yes/No','적용 여부와 Confirmed 상태를 모두 확인'),
            col('RuleSetVersion','기준 버전'),col('Note','설명','Multiple lines of text')],
        'rows':[
            {'Title':'PRODUCT_GI','RuleType':'Product','ProductFamily':'GI','SourceStandard':'ASTM A653/A653M',
             'SourceValue':'PRIME GALVANISED STEEL SHEET IN COILS','TargetValue':'GI','Status':'Confirmed','IsActive':True,
             'RuleSetVersion':'EXAMPLE_V1','Note':'제품군 분류 대응. 세부 강종·도금 공급 가능 확정을 의미하지 않음'},
            {'Title':'GRADE_SS37','RuleType':'Grade','ProductFamily':'GI','SourceStandard':'ASTM A653/A653M',
             'SourceValue':'SS37','TargetValue':None,'Status':'NeedsConfirmation','IsActive':False,
             'RuleSetVersion':'EXAMPLE_V1','Note':'확인한 GI 카탈로그 표에서 동등 강종을 확정하지 못함'},
            {'Title':'COATING_G60','RuleType':'Coating','ProductFamily':'GI','SourceStandard':'ASTM A653/A653M',
             'SourceValue':'G60','TargetValue':None,'Status':'NeedsConfirmation','IsActive':False,
             'RuleSetVersion':'EXAMPLE_V1','Note':'단위 환산만으로 현대제철 도금 등급을 확정하지 않음'}]
    },{
        'list_name':'SupplyRanges','display_name':'공급 범위 기준','grain':'제품군·강종 그룹·공장의 선택된 한 범위당 한 행',
        'columns':[
            col('Title','범위 규칙 키'),col('ProductFamily','제품군'),col('Plant','공장'),col('CatalogGradeGroup','카탈로그 강종 그룹'),
            col('ThicknessMinMm','두께 하한(mm)','Number'),col('ThicknessMinInclusive','두께 하한 포함','Yes/No'),
            col('ThicknessMaxMm','두께 상한(mm)','Number'),col('ThicknessMaxInclusive','두께 상한 포함','Yes/No'),
            col('WidthMinMm','폭 하한(mm)','Number'),col('WidthMinInclusive','폭 하한 포함','Yes/No'),
            col('WidthMaxMm','폭 상한(mm)','Number'),col('WidthMaxInclusive','폭 상한 포함','Yes/No'),
            col('Status','검토 상태','Choice','ExampleOnly / Confirmed'),col('IsActive','적용 여부','Yes/No'),
            col('RuleSetVersion','기준 버전'),col('Note','설명','Multiple lines of text')],
        'rows':[
            {'Title':'RANGE_SC_GI_01','ProductFamily':'GI','Plant':'순천공장','CatalogGradeGroup':'S310~S570M',
             'ThicknessMinMm':0.8,'ThicknessMinInclusive':True,'ThicknessMaxMm':2.3,'ThicknessMaxInclusive':False,
             'WidthMinMm':1040,'WidthMinInclusive':False,'WidthMaxMm':1250,'WidthMaxInclusive':True,
             'Status':'ExampleOnly','IsActive':False,'RuleSetVersion':'EXAMPLE_V1',
             'Note':'확인한 차트 일부 구간. SS37이 이 그룹에 속한다는 동등성은 미확정. 바로 운영 적용하지 않는 예시 행'},
            {'Title':'RANGE_DJ_GI_01','ProductFamily':'GI','Plant':'당진 1공장 #2 CGL','CatalogGradeGroup':'S330~S440',
             'ThicknessMinMm':0.35,'ThicknessMinInclusive':True,'ThicknessMaxMm':1.3,'ThicknessMaxInclusive':False,
             'WidthMinMm':775,'WidthMinInclusive':True,'WidthMaxMm':1275,'WidthMaxInclusive':False,
             'Status':'ExampleOnly','IsActive':False,'RuleSetVersion':'EXAMPLE_V1',
             'Note':'확인한 차트 일부 구간. SS37과의 동등성 미확정. 품목·강종·공장과 범위를 함께 비교해야 함'}]
    },{
        'list_name':'CellMappings','display_name':'엑셀 셀 매핑','grain':'템플릿 버전의 입력 항목 한 개당 한 행',
        'columns':[
            col('Title','매핑 키'),col('Department','부서','Choice'),col('TemplateKey','템플릿 키'),
            col('TemplateVersion','템플릿 버전'),col('SourceField','읽을 추출값 열'),col('WorksheetName','엑셀 시트명'),
            col('CellAddress','엑셀 셀 주소'),col('DataType','입력 자료형','Choice','Number / Text'),
            col('Required','필수 항목','Yes/No'),col('MissingValuePolicy','미확정·누락 처리','Choice','ReviewRequired'),
            col('IsActive','적용 여부','Yes/No'),col('Note','설명','Multiple lines of text')],
        'rows':[]
    }]
    for idx,(field,cell,datatype) in enumerate([('ThicknessMm','F10','Number'),('WidthMm','G10','Number'),
                                                ('HyundaiGrade','D10','Text'),('HyundaiCoating','E10','Text'),
                                                ('CoilInnerDiameterMm','H10','Number'),('QuantityT','I10','Number')],1):
        lists[3]['rows'].append({'Title':f'PUR_SPEC_{idx:02}','Department':'구매팀','TemplateKey':'PUR_SPEC','TemplateVersion':'1.0',
            'SourceField':field,'WorksheetName':'제품정보','CellAddress':cell,'DataType':datatype,'Required':True,
            'MissingValuePolicy':'ReviewRequired','IsActive':False,
            'Note':'시트·셀·템플릿은 설명용 예시. 실제 부서 양식 확인 후 적용 여부를 변경'})
    dataset={'description':'SharePoint 구성 예시. 실제 사이트에 생성·등록하지 않음. 페이지·좌표는 포함하지 않음.',
             'lists':lists,'template_library_example':{'library':'Templates','TemplateKey':'PUR_SPEC','Department':'구매팀',
                'Version':'1.0','Filename':'구매팀_제품검토서.xlsx','ExistsInThisSession':False,
                'Note':'실제 부서 템플릿을 문서 라이브러리에 올리고 TemplateKey로 연결하는 예시'}}
    (ROOT/'sharepoint_lists_example.json').write_text(json.dumps(dataset,ensure_ascii=False,indent=2),encoding='utf-8')
    workbook=Workbook();intro=workbook.active;intro.title='안내'
    intro.append(['구분','내용'])
    for row in [
        ['용도','SharePoint 목록별 열과 예시 행. 실제 SharePoint 사이트에는 생성하지 않았습니다.'],
        ['추출결과','발주서 DEMO001의 한 품목에서 추출한 값입니다. 여러 품목이면 품목별로 행을 늘립니다.'],
        ['대응기준','SS37·G60의 대응값은 미확정이므로 빈칸과 적용 안 함으로 남겼습니다.'],
        ['공급범위','두 차트의 일부 구간만 예시로 등록했습니다. SS37 동등성 미확정이며 적용 안 함입니다.'],
        ['셀매핑','구매팀·제품정보 시트·F10 등의 위치는 가상 예시입니다. 실제 양식 확인 전 적용하지 않습니다.'],
        ['가져오기','SharePoint 목록 만들기의 Excel 가져오기에서 해당 이름의 표를 선택합니다. 목록별로 각각 가져옵니다.'],
        ['열 형식','Excel 가져오기가 모든 열 형식을 완성하지는 않습니다. 열정의 시트대로 숫자·선택·예/아니오 등을 확인·설정합니다.'],
        ['필드 이름','영문 헤더로 생성하면 PA에서 참조할 내부 이름이 안정적입니다. 이후 표시 이름을 한국어로 바꿀 수 있습니다.'],
        ['템플릿','Excel 템플릿 파일은 목록의 텍스트 칸이 아니라 문서 라이브러리에 저장합니다.'],
        ['페이지 정보','이 예시 목록에는 페이지 정보가 없습니다. 필요할 때 원문 확인용 기록을 별도로 연결합니다.']]:intro.append(row)
    names={'ExtractedItems':'추출결과','MappingRules':'대응기준','SupplyRanges':'공급범위','CellMappings':'셀매핑'}
    definitions=workbook.create_sheet('열정의');definitions.append(['목록','영문 열 이름','한국어 표시 이름','권장 SharePoint 형식','설명'])
    for spec in lists:
        headers=[c['name'] for c in spec['columns']]
        for row in spec['rows']:assert set(row)==set(headers),(spec['list_name'],set(row)^set(headers))
        sheet=workbook.create_sheet(names[spec['list_name']]);sheet.append(headers)
        for row in spec['rows']:sheet.append([row.get(h) for h in headers])
        table=Table(displayName=spec['list_name'],ref=f'A1:{get_column_letter(len(headers))}{sheet.max_row}')
        table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True)
        sheet.add_table(table)
        for c in spec['columns']:definitions.append([spec['list_name'],c['name'],c['display_name'],c['recommended_sharepoint_type'],c['note']])
        with (ROOT/(spec['list_name']+'.csv')).open('w',encoding='utf-8-sig',newline='') as stream:
            writer=csv.DictWriter(stream,fieldnames=headers);writer.writeheader();writer.writerows(spec['rows'])
        for j,c in enumerate(spec['columns'],1):
            sheet.column_dimensions[get_column_letter(j)].width=24 if c['name']!='Note' else 70
            if c['recommended_sharepoint_type']=='Number':
                for row in sheet.iter_rows(min_row=2,min_col=j,max_col=j):
                    for cell in row:cell.number_format='0.###'
    for sheet in workbook:
        sheet.freeze_panes='A2'
        sheet.row_dimensions[1].height=30
        for c in sheet[1]:c.fill=PatternFill('solid',fgColor='1F4E78');c.font=Font(color='FFFFFF',bold=True);c.alignment=Alignment(wrap_text=True,vertical='center')
        for row in sheet.iter_rows(min_row=2):
            for c in row:c.alignment=Alignment(vertical='top',wrap_text=True)
    intro.column_dimensions['A'].width=18;intro.column_dimensions['B'].width=110
    for i in range(2,intro.max_row+1):intro.row_dimensions[i].height=34
    for colletter,width in [('A',22),('B',30),('C',30),('D',26),('E',100)]:definitions.column_dimensions[colletter].width=width
    path=ROOT/'SharePoint_수치기준표_구성예시.xlsx';workbook.save(path)
    reopened=load_workbook(path,data_only=False)
    assert reopened['추출결과']['I2'].value==0.85
    assert isinstance(reopened['추출결과']['G2'].value,str) and reopened['추출결과']['G2'].value=='SS37'
    assert reopened['대응기준']['F3'].value is None and reopened['대응기준']['H3'].value is False
    assert sum(len(s.tables) for s in reopened)==4
    assert all(r['IsActive'] is False for spec in lists[2:] for r in spec['rows'])
    print('WORKBOOK_VALID',path,'TABLES',4,'DATA_ROWS',sum(len(s['rows']) for s in lists))


if __name__=='__main__':main()
