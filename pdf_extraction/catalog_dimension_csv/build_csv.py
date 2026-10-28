#!/usr/bin/env python3
"""Create SharePoint-ready CSV reference lists from selected checked tables.

No tenant writes or app changes. Existing PDFs and OCR outputs are read-only.
Primary reference rows and unverified OCR review rows are exported separately.
"""

import os
import collections
import csv
from decimal import Decimal as D
import hashlib
import json
import math
from pathlib import Path
import re
import shutil
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

from curated_tables import TABLES
from plate_ranges import MATRICES,THICKNESS_BANDS,WIDTH_BANDS

ROOT=Path(__file__).resolve().parent
OUT=ROOT/'deliverables'
OCR=ROOT.parent/'results'
OUT.mkdir(exist_ok=True)
EXPORTS={}
AUDIT=[]


def num(value):
    if value in [None,'','-']:return None
    d=D(str(value).replace('*',''))
    return int(d) if d==d.to_integral_value() else float(d)


def canon(value):
    if value is None:return ''
    if isinstance(value,bool):return 'TRUE' if value else 'FALSE'
    if isinstance(value,(int,float,D)):
        s=format(D(str(value)).normalize(),'f')
        return '0' if s=='-0' else s
    return str(value)


def export(name,rows,headers=None):
    if headers is None:
        headers=list(dict.fromkeys(k for row in rows for k in row))
    p=OUT/name
    with p.open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fieldnames=headers,extrasaction='raise');w.writeheader()
        for row in rows:w.writerow({h:canon(row.get(h)) for h in headers})
    EXPORTS[name]={'rows':rows,'headers':headers}


def audit(key,doc,page,row,comment='',bbox=None):
    meta=json.loads((OCR/doc/'pages'/f'{page:04}.json').read_text())['source']
    AUDIT.append({'Title':f'AUDIT-{len(AUDIT)+1:06}','REFERENCE_KEY':key,'SOURCE_DOCUMENT':meta['filename'],
                  'SOURCE_PART':doc,'PDF_PAGE':page,'PRINTED_PAGE_CANDIDATE':meta.get('printed_page_label_candidate'),
                  'TABLE_ROW_OR_CELL':row,'BBOX_PT':json.dumps(bbox) if bbox else '',
                  'REVIEW_NOTE':comment})


EXACT_COLUMNS='''Title CATALOG_ITEM_NAME TARGET_ITEM_NAME TARGET_ITEM_CODE PRODUCT_TYPE MATERIAL_CODE STANDARD_SPEC GRADE SIZE_RAW THICKNESS WIDTH LENGTH INSIDE_DIAMETER OUTSIDE_DIAMETER WALL_THICKNESS SCHEDULE THICKNESS_TOLERANCE WIDTH_TOLERANCE SECTION_HEIGHT_MM LEG_A_MM LEG_B_MM LEG_A_THICKNESS_MM LEG_B_THICKNESS_MM WEB_THICKNESS_MM FLANGE_THICKNESS_MM CORNER_RADIUS_MM SECOND_RADIUS_MM NOMINAL_DIAMETER_MM SQUARE_SIDE_MM UNIT_WEIGHT_KG_M LB_PER_FT DIAMETER_BASIS DIMENSION_UNIT DIMENSION_SERIES DIMENSION_STANDARD MATCH_FIELDS MATCH_KEY SUPPLY_NOTE TRANSCRIPTION_STATUS ALLOW_DIMENSION_LOOKUP ALLOW_MATERIAL_CODE_REPLACEMENT RECORD_VERSION SOURCE_OCCURRENCES'''.split()


def build_exact():
    records={};profiles={};mass_warnings=[]
    for table in TABLES:
        family=table['family'];series=table['series']
        for index,values in enumerate(table['rows'],1):
            v=dict(zip(table['columns'],values));star=any('*' in s for s in values)
            row={k:None for k in EXACT_COLUMNS}
            row.update(PRODUCT_TYPE=family,DIMENSION_SERIES=series,DIMENSION_STANDARD=table['dimension_standard'],
                       DIMENSION_UNIT='mm',UNIT_WEIGHT_KG_M=num(v['W']),TRANSCRIPTION_STATUS='VISUALLY_CHECKED_DIMENSION_FIELDS',
                       ALLOW_DIMENSION_LOOKUP=True,ALLOW_MATERIAL_CODE_REPLACEMENT=False,RECORD_VERSION='V1',SOURCE_OCCURRENCES=1)
            nominal=v.get('NOMINAL','').replace('*','')
            row['SIZE_RAW']=nominal
            required=['PRODUCT_TYPE','DIMENSION_SERIES']
            if 'H' in v:
                row['SECTION_HEIGHT_MM']=num(v['H']);row['WIDTH']=num(v['B'])
            if 'T1' in v:
                if family=='INVERTED_ANGLE':
                    row['LEG_A_MM']=num(v['H']);row['LEG_B_MM']=num(v['B'])
                    row['LEG_A_THICKNESS_MM']=num(v['T1']);row['LEG_B_THICKNESS_MM']=num(v['T2'])
                    required+=['LEG_A_MM','LEG_B_MM','LEG_A_THICKNESS_MM','LEG_B_THICKNESS_MM']
                else:
                    row['WEB_THICKNESS_MM']=num(v['T1']);row['FLANGE_THICKNESS_MM']=num(v['T2'])
                    required+=['SECTION_HEIGHT_MM','WIDTH','WEB_THICKNESS_MM','FLANGE_THICKNESS_MM']
                shape='x'.join(canon(row[k]) for k in (['LEG_A_MM','LEG_B_MM','LEG_A_THICKNESS_MM','LEG_B_THICKNESS_MM'] if family=='INVERTED_ANGLE' else ['SECTION_HEIGHT_MM','WIDTH','WEB_THICKNESS_MM','FLANGE_THICKNESS_MM']))
            elif family in ['EQUAL_ANGLE','UNEQUAL_ANGLE']:
                a,b=[num(s) for s in nominal.split('x')]
                row.update(LEG_A_MM=a,LEG_B_MM=b,THICKNESS=num(v['T']),WIDTH=b)
                required+=['LEG_A_MM','LEG_B_MM','THICKNESS'];shape=f'{canon(a)}x{canon(b)}x{canon(row["THICKNESS"])}'
            elif family=='SHEET_PILE':
                row['THICKNESS']=num(v['T']);required+=['SECTION_HEIGHT_MM','WIDTH','THICKNESS']
                shape='x'.join(canon(row[k]) for k in ['SECTION_HEIGHT_MM','WIDTH','THICKNESS'])
            elif family in ['REBAR','THREAD_BAR']:
                if 'DIA' in v:
                    row['NOMINAL_DIAMETER_MM']=num(v['DIA']);row['DIAMETER_BASIS']='NOMINAL_EQUIVALENT_NOT_OUTER_RIB_DIAMETER'
                if family=='REBAR' and series=='BS_B500B':row['GRADE']='B500B'
                row['STANDARD_SPEC']=table['dimension_standard']
                row['LB_PER_FT']=num(v.get('LB_FT'))
                required+=['SIZE_RAW'];shape=nominal
            elif family in ['ROUND_BAR','WIRE_ROD']:
                diameter=num(v['DIA']);row.update(OUTSIDE_DIAMETER=diameter,NOMINAL_DIAMETER_MM=diameter,
                                                 DIAMETER_BASIS='CATALOG_NOMINAL_SOLID_ROUND_DIAMETER')
                row['SIZE_RAW']=f'Dia.{canon(diameter)}';required+=['OUTSIDE_DIAMETER'];shape=f'Dia.{canon(diameter)}'
            elif family=='SQUARE_BAR':
                side=num(v['SIDE']);row['SQUARE_SIDE_MM']=side;row['WIDTH']=side
                row['SIZE_RAW']=f'{canon(side)}x{canon(side)}';required+=['SQUARE_SIDE_MM'];shape=row['SIZE_RAW']
            else:raise ValueError(family)
            row['CORNER_RADIUS_MM']=num(v.get('R'));row['SECOND_RADIUS_MM']=num(v.get('R2'))
            row['CATALOG_ITEM_NAME']=f'{family} {shape} [{series}]'
            notes=[]
            if table['note']:notes.append(table['note'])
            if star:
                notes.append('THIS_ROW_HAS_A_SOURCE_ASTERISK; CHECK_SOURCE_RESTRICTION')
            if family=='JUNIOR_BEAM':notes.append('SPECIAL_ORDER_ALL_ROWS')
            row['SUPPLY_NOTE']='; '.join(notes)
            row['MATCH_FIELDS']='|'.join(required)
            row['MATCH_KEY']='|'.join(f'{k}={canon(row[k])}' for k in required)
            assert len(row['MATCH_KEY'])<=255,row['MATCH_KEY']
            key='HSD-'+hashlib.sha256(row['MATCH_KEY'].encode()).hexdigest()[:16].upper()
            row['Title']=key
            if key in records:
                old=records[key]
                assert old['UNIT_WEIGHT_KG_M']==row['UNIT_WEIGHT_KG_M'],('conflicting_duplicate',old,row)
                old['SOURCE_OCCURRENCES']+=1
            else:records[key]=row
            audit(key,table['doc'],table['page'],index,table['note'])
            profiles[(family,series)]=required
            # Independent dimensional sanity checks identify transcription mistakes.
            expected=None
            if family in ['ROUND_BAR','WIRE_ROD']:
                expected=math.pi*row['OUTSIDE_DIAMETER']**2/4*.00785
            elif family=='SQUARE_BAR':expected=row['SQUARE_SIDE_MM']**2*.00785
            elif family in ['H_SECTION','STEEL_H_PILE','JUNIOR_BEAM']:
                h,b,tw,tf,r=[row[k] or 0 for k in ['SECTION_HEIGHT_MM','WIDTH','WEB_THICKNESS_MM','FLANGE_THICKNESS_MM','CORNER_RADIUS_MM']]
                expected=(2*b*tf+(h-2*tf)*tw+(4-math.pi)*r*r)*.00785
            if expected is not None and abs(expected-row['UNIT_WEIGHT_KG_M'])>max(.15,expected*.025):
                mass_warnings.append({'REFERENCE_KEY':key,'FAMILY':family,'CATALOG_KG_M':row['UNIT_WEIGHT_KG_M'],'GEOMETRY_ESTIMATE_KG_M':round(expected,6)})
    rows=list(records.values())
    assert len(rows)==498
    assert len({r['MATCH_KEY'] for r in rows})==len(rows)
    assert all(r['LENGTH'] is None for r in rows)
    assert all(r['WALL_THICKNESS'] is None and r['SCHEDULE'] is None for r in rows)
    export('CatalogExactDimensions.csv',rows,EXACT_COLUMNS)
    profile_rows=[{'Title':f'{family}_{series}','PRODUCT_TYPE':family,'DIMENSION_SERIES':series,
                   'MATCH_MODE':'EXACT','REQUIRED_FIELDS':'|'.join(fields),'UNIT':'mm',
                   'MISSING_FIELD_POLICY':'REVIEW_REQUIRED','MULTIPLE_MATCH_POLICY':'REVIEW_REQUIRED',
                   'TARGET_POLICY':'Company TARGET_ITEM_NAME/TARGET_ITEM_CODE must be supplied and accepted separately'}
                  for (family,series),fields in profiles.items()]
    for product,required in [('GI','PRODUCT_TYPE|CATALOG_GRADE_GROUP|THICKNESS|WIDTH'),
                             ('STEEL_PLATE','PRODUCT_TYPE|PROCESS|THICKNESS|WIDTH|LENGTH')]:
        profile_rows.append({'Title':product+'_RANGE','PRODUCT_TYPE':product,'DIMENSION_SERIES':'CATALOG_SUPPLY_RANGE',
                             'MATCH_MODE':'RANGE','REQUIRED_FIELDS':required,'UNIT':'mm',
                             'MISSING_FIELD_POLICY':'REVIEW_REQUIRED','MULTIPLE_MATCH_POLICY':'REVIEW_REQUIRED',
                             'TARGET_POLICY':'Apply notes, edge/transport/grade limits; geometric fit is not material equivalence'})
    export('CatalogMatchProfiles.csv',profile_rows)
    export('DimensionSanityReview.csv',mass_warnings,['REFERENCE_KEY','FAMILY','CATALOG_KG_M','GEOMETRY_ESTIMATE_KG_M'])
    return rows,mass_warnings


RANGE_COLS='''Title CATALOG_ITEM_NAME TARGET_ITEM_NAME TARGET_ITEM_CODE PRODUCT_TYPE PROCESS CATALOG_GRADE_GROUP GRADE MATERIAL_CODE THICKNESS_MIN THICKNESS_MIN_INCLUSIVE THICKNESS_MAX THICKNESS_MAX_INCLUSIVE WIDTH_MIN WIDTH_MIN_INCLUSIVE WIDTH_MAX WIDTH_MAX_INCLUSIVE LENGTH_MIN LENGTH_MIN_INCLUSIVE TABLE_LENGTH_MAX_M TABLE_LENGTH_MAX_MM GENERAL_THICK_PLATE_LENGTH_CAP_MM LENGTH_MAX LENGTH_MAX_INCLUSIVE MILL_EDGE_WIDTH_MAX_MM SLIT_EDGE_WIDTH_MAX_MM LAND_TRANSPORT_LENGTH_MAX_MM SOURCE_CELL_STATUS MATCH_MODE DIMENSION_UNIT CONDITIONS_TRANSCRIBED ALLOW_DIMENSION_LOOKUP ALLOW_MATERIAL_CODE_REPLACEMENT NOTES RECORD_VERSION'''.split()


def build_ranges():
    rows=[]
    for page,matrix in MATRICES.items():
        for ri,(tlo,thi) in enumerate(THICKNESS_BANDS):
            for ci,(wlo,whi) in enumerate(WIDTH_BANDS[page]):
                text=matrix[ri][ci];m=num(text);length=num(D(text)*1000) if text!='-' else None
                key=f'HSPR-TMCP-{page}-{ri+1:02}-{ci+1:02}'
                if length is not None:status='NUMERIC_MAXIMUM'
                elif page==3:status='NOT_AVAILABLE'
                elif tlo>=80 and ci<6:status='CONSULTATION_REQUIRED_300MM'
                else:status='NO_NUMERIC_MAXIMUM_REQUIRES_CONFIRMATION'
                cap=13000 if tlo>=70 else None
                effective=min(length,cap) if length is not None and cap is not None else length
                row={k:None for k in RANGE_COLS}
                notes=['Source width bands are integer-mm bins; do not assume a fractional width in the gap between bins is covered.',
                       'Range transcription only. Material grade, heat treatment, high-grade production and order acceptance require additional conditions.',
                       'Maximum width for heat-treated product: mill edge 4700mm, slit edge 4750mm.',
                       'Maximum land-transport length: 21500mm. Apply only where that transport constraint is applicable.',
                       'Source minimum product length is 3m; all exported dimensions are mm.']
                if tlo>=70:notes.append('Source general note T>=70mm: max length 13m. Effective LENGTH_MAX is min(table maximum,13000); printed table value retained separately.')
                if thi>75 and wlo<=2100:notes.append('Some marine-structure grades with T>75mm and W<=2100mm require discussion; confirm grade-specific applicability.')
                if status!='NUMERIC_MAXIMUM':notes.append('No usable numeric maximum: do not interpret blank as zero or as unrestricted supply.')
                row.update(Title=key,CATALOG_ITEM_NAME=f'TMCP plate t[{tlo},{thi}) w[{wlo},{whi}]',
                           PRODUCT_TYPE='STEEL_PLATE',PROCESS='TMCP',THICKNESS_MIN=tlo,THICKNESS_MIN_INCLUSIVE=True,
                           THICKNESS_MAX=thi,THICKNESS_MAX_INCLUSIVE=False,WIDTH_MIN=wlo,WIDTH_MIN_INCLUSIVE=True,
                           WIDTH_MAX=whi,WIDTH_MAX_INCLUSIVE=True,LENGTH_MIN=3000,LENGTH_MIN_INCLUSIVE=True,
                           TABLE_LENGTH_MAX_M=m,TABLE_LENGTH_MAX_MM=length,GENERAL_THICK_PLATE_LENGTH_CAP_MM=cap,
                           LENGTH_MAX=effective,LENGTH_MAX_INCLUSIVE=True,MILL_EDGE_WIDTH_MAX_MM=4700,
                           SLIT_EDGE_WIDTH_MAX_MM=4750,LAND_TRANSPORT_LENGTH_MAX_MM=21500,SOURCE_CELL_STATUS=status,
                           MATCH_MODE='RANGE',DIMENSION_UNIT='mm',CONDITIONS_TRANSCRIBED=True,
                           ALLOW_DIMENSION_LOOKUP=length is not None,ALLOW_MATERIAL_CODE_REPLACEMENT=False,
                           NOTES=' '.join(notes),RECORD_VERSION='V1')
                rows.append(row);audit(key,'pt1_part2',page,f'thickness_row={ri+1};width_col={ci+1}',status)
    for key,plant,group,tlo,thi,wlo,whi,wli,wui,page in [
        ('HSPR-GI-SC-01','순천공장','S310~S570M',.8,2.3,1040,1250,False,True,90),
        ('HSPR-GI-DJ-01','당진 1공장 #2 CGL','S330~S440',.35,1.3,775,1275,True,False,105)]:
        row={k:None for k in RANGE_COLS}
        row.update(Title=key,CATALOG_ITEM_NAME=f'GI {group} selected band ({plant})',PRODUCT_TYPE='GI',
                   CATALOG_GRADE_GROUP=group,THICKNESS_MIN=tlo,THICKNESS_MIN_INCLUSIVE=True,
                   THICKNESS_MAX=thi,THICKNESS_MAX_INCLUSIVE=False,WIDTH_MIN=wlo,WIDTH_MIN_INCLUSIVE=wli,
                   WIDTH_MAX=whi,WIDTH_MAX_INCLUSIVE=wui,SOURCE_CELL_STATUS='SELECTED_BLUE_INTERIOR_REGION',
                   MATCH_MODE='RANGE',DIMENSION_UNIT='mm',CONDITIONS_TRANSCRIBED=True,ALLOW_DIMENSION_LOOKUP=True,
                   ALLOW_MATERIAL_CODE_REPLACEMENT=False,RECORD_VERSION='V1',
                   NOTES=f'{plant}. Selected interior portion only, not the complete supply chart. The catalog group is not an accepted mapping from ASTM SS37. Coil has no fixed LENGTH in this rule. All grade/coating/process requirements remain separate.')
        rows.append(row);audit(key,'pt1_part1',page,'selected_chart_band','Conditional geometric band, no grade-equivalence approval')
    assert len(rows)==562
    assert sum(r['TABLE_LENGTH_MAX_MM'] is not None for r in rows)==469
    assert all(r['LENGTH_MAX']<=13000 for r in rows if r['PRODUCT_TYPE']=='STEEL_PLATE' and r['THICKNESS_MIN']>=70 and r['LENGTH_MAX'] is not None)
    export('CatalogSupplyRanges.csv',rows,RANGE_COLS)
    return rows


def build_field_definitions():
    source_fields='''ITEM_NO DESCRIPTION PRODUCT_TYPE MATERIAL_CODE STANDARD_SPEC GRADE SIZE_RAW THICKNESS WIDTH LENGTH INSIDE_DIAMETER OUTSIDE_DIAMETER WALL_THICKNESS SCHEDULE ORDER_QTY QTY_UOM UNIT_PRICE PRICE_UOM LINE_AMOUNT EDGE SKIN_PASS COIL_WEIGHT_MIN COIL_WEIGHT_MAX COATING_WEIGHT SURFACE THICKNESS_TOLERANCE WIDTH_TOLERANCE WEIGHT_PER_SHEET UT_TEST HEAT_NO END_TYPE PIPE_COATING PCS_PER_BUNDLE BUNDLE_QTY TOTAL_PCS LB_PER_FT'''.split()
    assert len(source_fields)==36
    direct={'SIZE_RAW','THICKNESS','WIDTH','LENGTH','INSIDE_DIAMETER','OUTSIDE_DIAMETER','WALL_THICKNESS'}
    tolerance={'THICKNESS_TOLERANCE','WIDTH_TOLERANCE'}
    categorical={'SCHEDULE'}
    qualifiers={'PRODUCT_TYPE','MATERIAL_CODE','STANDARD_SPEC','GRADE','EDGE','SKIN_PASS','SURFACE','END_TYPE','PIPE_COATING'}
    weight={'COIL_WEIGHT_MIN','COIL_WEIGHT_MAX','COATING_WEIGHT','WEIGHT_PER_SHEET','LB_PER_FT'}
    textfields=direct&{'SIZE_RAW'}|tolerance|categorical|qualifiers|{'ITEM_NO','DESCRIPTION','QTY_UOM','PRICE_UOM','UT_TEST','HEAT_NO'}
    labels={'ITEM_NO':'문서 품목 순번','DESCRIPTION':'품목 설명','PRODUCT_TYPE':'제품 형태','MATERIAL_CODE':'재질·자재 코드',
            'STANDARD_SPEC':'적용 규격','GRADE':'강종','SIZE_RAW':'원문 치수 문자열','THICKNESS':'두께','WIDTH':'폭','LENGTH':'제품 길이',
            'INSIDE_DIAMETER':'내경','OUTSIDE_DIAMETER':'외경','WALL_THICKNESS':'관벽 두께','SCHEDULE':'관 스케줄',
            'THICKNESS_TOLERANCE':'두께 공차 원문','WIDTH_TOLERANCE':'폭 공차 원문','COIL_WEIGHT_MIN':'코일 최소중량',
            'COIL_WEIGHT_MAX':'코일 최대중량','COATING_WEIGHT':'도금 부착량','WEIGHT_PER_SHEET':'장당 중량','LB_PER_FT':'단위 길이당 중량'}
    notes={
        'SIZE_RAW':'원문을 보존. x/× 표기만으로 항상 두께·폭·길이 순서라고 가정하지 말 것.',
        'THICKNESS':'판·코일 두께 또는 단일 두께 형강 t. H형강의 t1/t2를 하나의 두께로 합치지 않음.',
        'WIDTH':'판·코일 폭 또는 명시된 단면 B. 제품 형태를 함께 비교.',
        'LENGTH':'제품 축방향 길이. H형강 H나 ㄱ형강 A를 길이로 입력하지 않음. 코일·단면표의 미기재 길이는 빈칸.',
        'INSIDE_DIAMETER':'코일 내경과 관 내경은 PRODUCT_TYPE을 구분. 발주서 508mm를 공통 카탈로그 규격으로 복제하지 않음.',
        'OUTSIDE_DIAMETER':'관 외경 또는 원형 봉·선재의 카탈로그 직경. 철근의 리브 포함 실제 외경과 공칭직경은 다름.',
        'WALL_THICKNESS':'관벽 두께. H형강 웹 두께와 같은 열로 치환하지 않음. 제공 카탈로그에서 관벽 표를 확정하지 못함.',
        'SCHEDULE':'SCH40/40S 등 범주형. 숫자 40을 mm로 취급하지 않음. 제공 카탈로그에서 스케줄 표 미확인.',
        'THICKNESS_TOLERANCE':'원문 텍스트 보존 후 부호·상하한·단위를 별도 수치 열로 분리. 임의 반올림 금지.',
        'WIDTH_TOLERANCE':'원문 텍스트 보존 후 상하한 mm 열로 분리. 미기재를 0으로 채우지 않음.',
        'COIL_WEIGHT_MIN':'기하 치수가 아닌 중량 조건. 정규화 기준 kg.',
        'COIL_WEIGHT_MAX':'기하 치수가 아닌 중량 조건. 정규화 기준 kg.',
        'COATING_WEIGHT':'도금 질량/면적. G60 같은 도금 명칭은 별도 COATING_DESIGNATION 텍스트 열로 보존.',
        'WEIGHT_PER_SHEET':'기하 치수가 아닌 장당 중량 kg. 단위 길이당 kg/m와 구분.',
        'LB_PER_FT':'원문 lb/ft 수치를 보존. kg/m로 환산할 경우 단위와 환산 근거를 별도 기록.',
        'STANDARD_SPEC':'재질·제품 규격. 치수표의 ASTM A6를 재질 ASTM A36과 같다고 취급하지 않음.',
        'MATERIAL_CODE':'OCR에서 읽은 원래 자재 코드는 보존. 회사의 최종 대응 코드는 TARGET_ITEM_CODE에 별도 등록.',
        'ITEM_NO':'통상 발주 행번호이며 회사 품목코드와 구분.',
        'SKIN_PASS':'YES/NO/UNKNOWN 선택형 권장. 미확인을 false로 바꾸지 않음.'}
    rows=[]
    for f in source_fields:
        group='직접 치수' if f in direct else '공차' if f in tolerance else '치수 관련 범주' if f in categorical else '제품·공정 조건' if f in qualifiers else '중량·도금 조건' if f in weight else '거래·수량·추적 정보'
        unit='mm' if f in direct-{'SIZE_RAW'} else 'kg' if f in {'COIL_WEIGHT_MIN','COIL_WEIGHT_MAX','WEIGHT_PER_SHEET'} else 'g/m²' if f=='COATING_WEIGHT' else 'lb/ft' if f=='LB_PER_FT' else ''
        rows.append({'Title':f,'DISPLAY_NAME':labels.get(f,f),'CATEGORY':group,
                     'SHAREPOINT_TYPE':'Choice' if f in ['PRODUCT_TYPE','SKIN_PASS','EDGE'] else 'Text' if f in textfields else 'Number',
                     'NORMALIZED_UNIT':unit,'USE_FOR_GEOMETRY_MATCH':f in direct-{'SIZE_RAW'},
                     'NULL_POLICY':'Keep blank; never infer zero or an unrestricted wildcard',
                     'NOTES':notes.get(f,'제품 조건에 따라 보조 필터로 사용' if f in qualifiers else '치수 조합 키에서 제외')})
    for field,label,unit,note in [
        ('SECTION_HEIGHT_MM','단면 높이 H','mm','제품 길이 LENGTH와 구분'),('LEG_A_MM','ㄱ형강 다리 A','mm','단면 형상 A'),
        ('LEG_B_MM','ㄱ형강 다리 B','mm','단면 형상 B'),('LEG_A_THICKNESS_MM','ㄱ형강 다리 A 두께','mm','부등두께 ㄱ형강 t1'),
        ('LEG_B_THICKNESS_MM','ㄱ형강 다리 B 두께','mm','부등두께 ㄱ형강 t2'),
        ('WEB_THICKNESS_MM','웹 두께 t1','mm','H/I/채널의 두께. 관벽 두께가 아님'),
        ('FLANGE_THICKNESS_MM','플랜지 두께 t2','mm','H/I/채널의 플랜지'),
        ('CORNER_RADIUS_MM','반경 r 또는 r1','mm','규격판본 구분과 함께 사용'),('SECOND_RADIUS_MM','반경 r2','mm','미기재 대시는 빈칸'),
        ('NOMINAL_DIAMETER_MM','공칭직경','mm','철근의 리브 포함 외경으로 치환하지 않음'),
        ('SQUARE_SIDE_MM','각봉 한 변','mm','각봉의 Dia 표기를 관 외경으로 해석하지 않음'),
        ('UNIT_WEIGHT_KG_M','단위 길이 중량','kg/m','장당 중량과 구분'),
        ('THICKNESS_TOLERANCE_PLUS_MM','두께 상향 공차','mm','원문 단위를 확인 후 정규화'),
        ('THICKNESS_TOLERANCE_MINUS_MM','두께 하향 공차','mm','원문 단위를 확인 후 정규화'),
        ('WIDTH_TOLERANCE_PLUS_MM','폭 상향 공차','mm','원문 단위를 확인 후 정규화'),
        ('WIDTH_TOLERANCE_MINUS_MM','폭 하향 공차','mm','원문 단위를 확인 후 정규화'),
        ('TARGET_ITEM_NAME','회사 최종 대응 항목명','','예: 실제 업무의 123. 제공 자료에 없어 현재 빈칸'),
        ('TARGET_ITEM_CODE','회사 최종 품목코드','','관리자가 확정해 등록하는 대응 코드'),
        ('COATING_DESIGNATION','도금 명칭','','G60 등의 원문 도금 분류; COATING_WEIGHT 숫자와 구분')]:
        rows.append({'Title':field,'DISPLAY_NAME':label,'CATEGORY':'추가 권장 열',
                     'SHAREPOINT_TYPE':'Number' if unit else 'Text','NORMALIZED_UNIT':unit,
                     'USE_FOR_GEOMETRY_MATCH':bool(unit=='mm'),'NULL_POLICY':'Keep blank when absent','NOTES':note})
    export('OCR_FieldDefinitions.csv',rows)
    return rows


def build_units_and_aliases():
    unitrows=[]
    for source,factor in [('mm','1'),('cm','10'),('m','1000'),('in','25.4'),('inch','25.4'),('ft','304.8')]:
        unitrows.append({'Title':'LENGTH_'+source,'QUANTITY_KIND':'Length','SOURCE_UNIT':source,'TARGET_UNIT':'mm',
                         'MULTIPLY_BY':num(factor),'ROUND_FOR_MATCH':False,'REQUIRES_EXPLICIT_SOURCE_UNIT':True,
                         'NOTES':'Never infer inch from a decimal value or a document language alone.'})
    for source,factor in [('kg','1'),('t','1000'),('lb','0.45359237')]:
        unitrows.append({'Title':'MASS_'+source,'QUANTITY_KIND':'Mass','SOURCE_UNIT':source,'TARGET_UNIT':'kg',
                         'MULTIPLY_BY':num(factor),'ROUND_FOR_MATCH':False,'REQUIRES_EXPLICIT_SOURCE_UNIT':True,
                         'NOTES':'TON must be resolved to metric/short/long ton from source context.'})
    unitrows.append({'Title':'LINE_MASS_LB_FT','QUANTITY_KIND':'MassPerLength','SOURCE_UNIT':'lb/ft','TARGET_UNIT':'kg/m',
                     'MULTIPLY_BY':float(D('.45359237')/D('.3048')),'ROUND_FOR_MATCH':False,'REQUIRES_EXPLICIT_SOURCE_UNIT':True,'NOTES':'Factor derived from exact pound and foot definitions.'})
    export('UnitNormalization.csv',unitrows)
    aliases=[]
    for family,names in {'H_SECTION':['H Section','H형강'],'EQUAL_ANGLE':['Equal Angle','등변 ㄱ형강'],
                         'UNEQUAL_ANGLE':['Unequal Angle','부등변 ㄱ형강'],'CHANNEL':['Channel','ㄷ형강'],
                         'REBAR':['Reinforcing Bar','철근'],'THREAD_BAR':['Thread Bar','나사형 철근'],
                         'ROUND_BAR':['Round Bar','원형 봉강'],'WIRE_ROD':['Wire Rod','선재'],
                         'SQUARE_BAR':['Square Bar','각봉'],'GI':['GI','Galvanized steel sheet in coils'],
                         'STEEL_PLATE':['Steel Plate','후판']}.items():
        for name in names:aliases.append({'Title':f'ALIAS-{len(aliases)+1:03}','SOURCE_PRODUCT_TYPE':name,'CANONICAL_PRODUCT_TYPE':family,
                                         'IS_CONFIRMED_PRODUCT_FAMILY_ALIAS':True,'NOTES':'Product form only; no grade equivalence or supply approval.'})
    aliases.append({'Title':'ALIAS-HR-AMBIGUOUS','SOURCE_PRODUCT_TYPE':'HR','CANONICAL_PRODUCT_TYPE':'',
                    'IS_CONFIRMED_PRODUCT_FAMILY_ALIAS':False,'NOTES':'Hot rolled alone does not establish coil vs plate; confirm form before matching.'})
    export('ProductTypeAliases.csv',aliases)


def build_current_input_review():
    p=Path(os.environ.get('WHATNOT_OCR_INPUT', 'private/ocr_input.json'))
    d=json.loads(p.read_text());rows=[]
    def val(x):return x.get('value') if isinstance(x,dict) else x
    for i,item in enumerate(d['PO_ITEM']['value'],1):
        rows.append({'Title':f'PO-{val(d["PO_NUMBER"])}-{i:03}','PURCHASE_ORDER':val(d['PO_NUMBER']),
                     'PRODUCT_TYPE':val(item.get('PRODUCT_TYPE')),'STANDARD_SPEC':val(item.get('STANDARD_SPEC')),
                     'SIZE_RAW':val(item.get('SIZE_RAW')),'THICKNESS':val(item.get('THICKNESS')),'WIDTH':val(item.get('WIDTH')),
                     'LENGTH':val(item.get('LENGTH')),'DIMENSION_UNIT':None,'MATCH_STATUS':'REVIEW_REQUIRED',
                     'NOTES':'THICKNESS and WIDTH fields are absent in this supplied result. SIZE_RAW contains 10 x 200 x 3000; source unit and separator are unconfirmed. Do not assume inches or invent a length. HR form needs coil/plate clarification.'})
    export('CurrentOCRInputReview.csv',rows)


def build_review_and_coverage(checked_pages):
    cells=[];coverage=[];counter=collections.Counter()
    for doc in ['pt1_part1','pt1_part2','pt2_part1','pt2_part2']:
        for path in sorted((OCR/doc/'pages').glob('*.json')):
            data=json.loads(path.read_text());page=data['source']['pdf_page_1based'];text=data['raw_text'];low=text.lower()
            kind='OTHER'
            if 'available sizes' in low or '공급 가능 규격' in text or (doc=='pt1_part1' and 78<=page<=120) or (doc=='pt1_part2' and 1<=page<=4):kind='SUPPLY_SIZE'
            if doc=='pt2_part1' and 18<=page<=101:kind='SECTION_DIMENSIONS_OR_PROPERTIES'
            if doc=='pt2_part1' and 102<=page<=119:kind='BAR_DIMENSIONS_OR_PACKAGING'
            if doc=='pt2_part2' and 27<=page<=46:kind='DIMENSION_OR_WEIGHT_TOLERANCE'
            if doc=='pt2_part2' and page in [47,48]:kind='BAR_DIMENSIONS'
            checked=(doc,page) in checked_pages
            coverage.append({'Title':f'{doc}_{page:04}','SOURCE_DOCUMENT':data['source']['filename'],'PDF_PAGE':page,
                             'CATEGORY':kind,'HAS_CURATED_DIMENSION_ROWS':checked,
                             'REVIEW_NEEDED_FOR_OTHER_DIMENSION_CONTENT':kind!='OTHER' and not checked,
                             'NOTE':'Curated rows cover declared dimension/weight fields only, not every property or note on a page.' if checked else ''})
            if kind=='OTHER' or checked:continue
            for table in data['table_grid_candidates']:
                for cell in table['cells']:
                    if not cell['ocr_text'].strip():continue
                    raw=cell['ocr_text'];escaped=raw.lstrip().startswith(('=','+','-','@'))
                    key=f'REVIEW-{len(cells)+1:06}'
                    cells.append({'Title':key,'CATEGORY':kind,'OCR_TEXT':("'"+raw) if escaped else raw,
                                  'TEXT_ESCAPED_FOR_SPREADSHEET':escaped,
                                  'GRID_ID':table['id'],'CELL_ID':cell['id'],
                                  'SAMPLED_RGB':','.join(str(x) for x in cell['sampled_background_rgb_0_255']),
                                  'TRANSCRIPTION_STATUS':'UNVERIFIED_OCR_ONLY','ALLOW_DIMENSION_LOOKUP':False,
                                  'NOTES':'Review-only. Table headers, units, grade context, merged cells and color meanings are not assigned.'})
                    # Do not repeat source JSON parsing for thousands of cells.
                    AUDIT.append({'Title':f'AUDIT-{len(AUDIT)+1:06}','REFERENCE_KEY':key,'SOURCE_DOCUMENT':data['source']['filename'],
                                  'SOURCE_PART':doc,'PDF_PAGE':page,'PRINTED_PAGE_CANDIDATE':data['source'].get('printed_page_label_candidate'),
                                  'TABLE_ROW_OR_CELL':cell['id'],'BBOX_PT':json.dumps(cell['bbox_pt']),
                                  'REVIEW_NOTE':'Unverified OCR cell, not an import-ready geometric record.'})
                    counter[kind]+=1
    assert len(coverage)==464
    export('CatalogReviewCells_DO_NOT_AUTOMATCH.csv',cells)
    export('CatalogCoverage.csv',coverage)
    return counter


def make_workbook():
    book=Workbook();intro=book.active;intro.title='안내';intro.append(['구분','내용'])
    for line in [
        ('실행용 목록','정확규격과 공급범위는 별도 목록으로 가져옵니다. 회사 최종 항목명·품목코드는 별도 확정이 필요합니다.'),
        ('제공 범위','원본 이미지와 대조한 21쪽의 단면·봉강·선재 치수 498행, TMCP 2쪽 560셀과 GI 선택 구간 2행.'),
        ('전체 검수 아님','464쪽 전체를 계산용 기준표로 검수한 결과는 아닙니다. 나머지 치수 관련 OCR 셀은 검토 전용 CSV에 있습니다.'),
        ('SharePoint 가져오기','CSV 가져오기를 지원하면 CSV를 선택합니다. Excel 가져오기에서는 이 파일의 해당 표를 선택합니다.'),
        ('열 형식','불러온 뒤 숫자·예/아니오·여러 줄 텍스트 열 형식을 확인합니다. 영문 열 이름은 OCR 키와 일치시켰습니다.'),
        ('내부 키','Title/MATCH_KEY는 이 데이터의 내부 참조키입니다. 현대제철 공식 품목코드가 아닙니다.'),
        ('회사 코드','TARGET_ITEM_NAME/TARGET_ITEM_CODE는 실제 회사 매핑값으로 채웁니다. 미확정 상태에서 자동 자재코드 치환을 켜지 않습니다.'),
        ('단면과 길이','H형강 H, ㄱ형강 A/B는 단면치수입니다. 축방향 LENGTH로 넣지 않았습니다.'),
        ('현재 OCR','DEMO002의 제공 JSON에는 THICKNESS/WIDTH가 없고 SIZE_RAW=10 x 200 x 3000만 있어 단위와 치수 분리가 필요합니다.'),
        ('미확인 제품','관 스케줄·관벽·관 내외경 표를 제공 카탈로그에서 확정하지 못했습니다. 빈칸은 0이 아닙니다.'),
        ('출처','페이지와 좌표는 CatalogSourceAudit.csv에 분리해 보관했습니다. 실행용 규격/범위 CSV에는 페이지를 넣지 않았습니다.')]:intro.append(line)
    selection=[('CatalogExactDimensions.csv','정확규격','ExactDimensions'),('CatalogSupplyRanges.csv','공급범위','SupplyRanges'),
               ('OCR_FieldDefinitions.csv','OCR열정의','FieldDefinitions'),('CatalogMatchProfiles.csv','제품별매칭조건','MatchProfiles'),
               ('UnitNormalization.csv','단위환산','UnitRules'),('ProductTypeAliases.csv','제품명표준화','ProductAliases'),
               ('CurrentOCRInputReview.csv','현재OCR검토','CurrentOCRReview')]
    for filename,sheetname,tablename in selection:
        content=EXPORTS[filename];sheet=book.create_sheet(sheetname);sheet.append(content['headers'])
        for row in content['rows']:
            values=[row.get(h) for h in content['headers']]
            # Add an apostrophe only to actual spreadsheet formula-like text.
            sheet.append(["'"+v if isinstance(v,str) and v.lstrip().startswith(('=','+','@')) else v for v in values])
        t=Table(displayName=tablename,ref=f'A1:{get_column_letter(sheet.max_column)}{sheet.max_row}')
        t.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True);sheet.add_table(t)
        for c in range(1,sheet.max_column+1):sheet.column_dimensions[get_column_letter(c)].width=24
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment=Alignment(vertical='top',wrap_text=True)
                if cell.data_type=='n':cell.number_format='0.#####'
    for sheet in book:
        sheet.freeze_panes='A2';sheet.row_dimensions[1].height=34
        for c in sheet[1]:c.font=Font(color='FFFFFF',bold=True);c.fill=PatternFill('solid',fgColor='1F4E78');c.alignment=Alignment(wrap_text=True,vertical='center')
    intro.column_dimensions['A'].width=24;intro.column_dimensions['B'].width=130
    for i in range(2,intro.max_row+1):intro.row_dimensions[i].height=38;intro.cell(i,2).alignment=Alignment(wrap_text=True,vertical='top')
    path=OUT/'SharePoint_CatalogDimensions.xlsx';book.save(path)
    reopened=load_workbook(path,data_only=False)
    assert reopened['정확규격'].max_row==499 and reopened['공급범위'].max_row==563
    assert sum(len(sheet.tables) for sheet in reopened)==7


def main():
    exact,warnings=build_exact();ranges=build_ranges();defs=build_field_definitions()
    build_units_and_aliases();build_current_input_review()
    checked={(t['doc'],t['page']) for t in TABLES}|{('pt1_part2',2),('pt1_part2',3),('pt1_part1',90),('pt1_part1',105)}
    review=build_review_and_coverage(checked)
    export('CatalogSourceAudit.csv',AUDIT)
    make_workbook()
    report={'input_OCR_fields_classified':36,'exact_dimension_rows':len(exact),'visually_transcribed_input_rows':500,
            'identical_printed_duplicate_rows_merged':2,'exact_families':dict(collections.Counter(r['PRODUCT_TYPE'] for r in exact)),
            'TMCP_cells_transcribed':560,'TMCP_numeric_length_cells':469,'GI_selected_interior_bands':2,
            'range_rows':len(ranges),'active_geometric_range_rows':sum(r['ALLOW_DIMENSION_LOOKUP'] for r in ranges),
            'review_cells_not_for_auto_match':sum(review.values()),'coverage_pages_scanned':464,
            'curated_exact_dimension_pages':21,'curated_TMCP_pages':2,'geometry_sanity_flags':warnings,
            'source_PDFs_or_existing_OCR_outputs_modified':False,'SharePoint_lists_created':False,
            'material_code_mapping_supplied':False,'entire_catalog_verified':False,
            'checks':{'csv_UTF8_BOM':True,'unique_exact_match_keys':True,'no_invented_product_lengths':True,
                      'no_invented_pipe_schedules':True,'source_code_fields_blank':True,'thick_plate_13m_note_applied':True}}
    (OUT/'validation_summary.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    # Round-trip read all generated CSVs, guarding against column shifts/newlines.
    for name,content in EXPORTS.items():
        assert (OUT/name).read_bytes().startswith(b'\xef\xbb\xbf')
        with (OUT/name).open(encoding='utf-8-sig',newline='') as stream:
            read=list(csv.DictReader(stream))
        assert len(read)==len(content['rows']) and all(None not in row for row in read),name
    print(json.dumps(report,ensure_ascii=False,indent=2))


if __name__=='__main__':main()
